"""Import product builds and their test reports from a test-report workbook.

Expected layout (one file, many sheets):
  * "Summary": header row 3, then one row per report
    (Report ID, Date, Submitted At, Operator, Product, Serial Number, Build Reference,
     Procedure, Overall Result, ..., Comments).
  * One sheet per report, named by Report ID: two header rows of key/value pairs, then
    a "Step | Criteria | Result | Comments | Sign Off" table whose section titles are
    rows with no criteria.

Each report becomes a serialised Unit on a completed work order (reports from the same
procedure on the same day share a work order), plus a TestReport with its steps.
"""
from datetime import datetime, timedelta, timezone as dt_timezone

from django.db import transaction
from django.utils import timezone

from .models import (Event, Product, ProductionTechnician, TestReport, TestStep, Unit, WorkOrder,
                     Workstation)


def _text(value):
    return "" if value is None else str(value).strip()


def _as_date(value):
    return value.date() if isinstance(value, datetime) else value


def _as_datetime(value, fallback_date):
    if isinstance(value, str) and value:
        try:
            return datetime.fromisoformat(value.replace("Z", "+00:00"))
        except ValueError:
            pass
    if isinstance(value, datetime):
        return value if value.tzinfo else value.replace(tzinfo=dt_timezone.utc)
    return datetime.combine(fallback_date, datetime.min.time(), tzinfo=dt_timezone.utc) + timedelta(hours=12)


def _match_technician(raw, aliases):
    """Map a free-text operator name to a ProductionTechnician, or None."""
    key = _text(raw).lower()
    if not key:
        return None
    wanted = (aliases or {}).get(key) or key.split()[0]
    for tech in ProductionTechnician.objects.all():
        if tech.name.lower() == wanted.lower():
            return tech
    return None


def _parse_steps(ws):
    steps, section, order = [], "", 0
    for row in ws.iter_rows(min_row=5, values_only=True):
        number, criteria, result, comments, sign_off = (list(row) + [None] * 5)[:5]
        if number is None and criteria is None:
            continue
        if criteria is None and result is None:  # section heading, e.g. "2. Visual Inspection"
            section = _text(number)
            continue
        number = _text(number)
        sec_no = section.split(".")[0]
        # A spreadsheet stores 3.10 as the number 3.1; restore it when it follows 3.9
        if steps and number == f"{sec_no}.1" and steps[-1]["section"] == section:
            number = f"{sec_no}.10"
        order += 1
        steps.append({"order": order, "section": section, "number": number, "criteria": _text(criteria),
                      "result": _text(result).upper(), "comments": _text(comments),
                      "sign_off": _text(sign_off)})
    return steps


@transaction.atomic
def import_test_reports(path, aliases=None, product_names=None, workstation="Test bay"):
    """Import the workbook at `path`. Returns the number of reports created."""
    import openpyxl  # imported here so the rest of the app doesn't need it

    wb = openpyxl.load_workbook(path, data_only=True)
    summary = wb["Summary"]
    rows = [r for r in summary.iter_rows(min_row=4, values_only=True) if r and r[0]]

    station, _ = Workstation.objects.get_or_create(name=workstation)
    numbers = [int(n) for n in WorkOrder.objects.values_list("number", flat=True) if n.isdigit()]
    next_number = max(numbers, default=3400) + 1
    work_orders = {}  # (procedure, date) -> WorkOrder
    created = 0

    for row in rows:
        row = tuple(row) + (None,) * (13 - len(row))
        report_id, date, submitted, operator, product_label, serial, build_ref, procedure, overall = row[:9]
        comments = row[12]
        if report_id not in wb.sheetnames or TestReport.objects.filter(report_id=report_id).exists():
            continue
        test_date = _as_date(date)
        procedure = _text(procedure)
        name = (product_names or {}).get(procedure) or _text(product_label) or procedure
        product, _ = Product.objects.get_or_create(code=procedure, defaults={"name": name})
        tech = _match_technician(operator, aliases)

        key = (procedure, test_date)
        if key not in work_orders:
            work_orders[key] = WorkOrder.objects.create(
                number=str(next_number), product=product, quantity=0, workstation=station,
                technician=tech, due_date=test_date, status=WorkOrder.COMPLETE, printed=True)
            next_number += 1
        order = work_orders[key]
        order.quantity += 1
        order.save()

        when = _as_datetime(submitted, test_date)
        passed = _text(overall).upper() == TestReport.PASS
        steps = _parse_steps(wb[report_id])
        failed = [s for s in steps if s["result"] == "FAIL"]
        unit = Unit.objects.create(
            work_order=order, serial=_text(serial), result=Unit.PASS if passed else Unit.REWORK,
            reason="" if passed else f"{len(failed)} test step(s) failed", first_pass=passed)
        Unit.objects.filter(pk=unit.pk).update(created_at=when)

        report = TestReport.objects.create(
            report_id=report_id, procedure=procedure, product=product, unit=unit, serial=_text(serial),
            build_reference=_text(build_ref), operator_name=_text(operator), technician=tech,
            test_date=test_date, submitted_at=when, overall_result=_text(overall).upper(),
            comments=_text(comments))
        TestStep.objects.bulk_create([TestStep(report=report, **s) for s in steps])

        for action, delta, detail in [("unit recorded", 0, f"{'pass' if passed else 'fail'} · test report {report_id}"),
                                      ("sent to QA", 5, ""), ("QA approved, completed", 30, "")]:
            event = Event.objects.create(work_order=order, unit=unit if delta == 0 else None,
                                         action=action, detail=detail)
            Event.objects.filter(pk=event.pk).update(timestamp=when + timedelta(minutes=delta))
        created += 1
    return created
